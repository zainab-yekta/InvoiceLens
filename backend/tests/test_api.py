import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import database
import main

client = TestClient(main.app)

GOOD_FIELDS = {
    "invoice_number": "RE-1", "date": "28.07.2025", "total_amount": "1.785,00",
    "tax_number": "12/345/67890", "vat_amount": "285,00", "vat_percent": "19", "currency": "eur",
}


@pytest.fixture(autouse=True)
def empty_database(tmp_path, monkeypatch):
    monkeypatch.setattr(database, "DB_NAME", str(tmp_path / "test.db"))
    monkeypatch.setattr(database, "UPLOAD_DIR", str(tmp_path / "uploads"))
    database.init_db()


def invoices():
    return client.get("/get_invoices").json()


def test_save_cleans_german_values_and_records_history():
    upload_id = database.store_upload(b"%PDF-1.4 test", ".pdf")
    response = client.post("/save_invoice", json={
        "fields": GOOD_FIELDS, "language": "de", "tags": ["consulting"],
        "upload_id": upload_id, "original_filename": "rechnung.pdf",
    })
    assert response.status_code == 200

    saved = invoices()["accepted"][0]
    assert saved["tags"] == "consulting"
    assert saved["total_amount"] == "1785.00"
    assert saved["issue_date"] == "2025-07-28"
    assert saved["currency"] == "EUR"

    history = invoices()["history"][0]
    assert history["status"] == "accepted"
    assert history["tags"] == "consulting"
    assert history["has_file"] is True

    file_response = client.get(f"/files/{history['id']}")
    assert file_response.status_code == 200
    assert file_response.content == b"%PDF-1.4 test"
    assert "rechnung.pdf" in file_response.headers["content-disposition"]


def test_duplicate_and_missing_fields():
    assert client.post("/save_invoice", json={"fields": {"invoice_number": "X"}}).status_code == 400
    assert client.post("/save_invoice", json={"fields": GOOD_FIELDS, "language": "de"}).status_code == 200
    assert client.post("/save_invoice", json={"fields": GOOD_FIELDS, "language": "de"}).status_code == 409


def test_reject_records_history_and_allows_several_unknown():
    for _ in range(2):
        response = client.post("/reject_invoice", json={"fields": {}, "reason": "No text could be read"})
        assert response.status_code == 200
    data = invoices()
    assert [r["invoice_number"] for r in data["rejected"]] == ["UNKNOWN", "UNKNOWN"]
    assert [h["status"] for h in data["history"]] == ["rejected", "rejected"]
    assert data["history"][0]["has_file"] is False


def test_file_link_rejects_made_up_names():
    database.add_history("rejected", "", {}, stored_filename="../../invoices.db")
    history_id = invoices()["history"][0]["id"]
    assert client.get(f"/files/{history_id}").status_code == 404


def test_edit_saved_invoice():
    client.post("/save_invoice", json={"fields": GOOD_FIELDS, "language": "de"})
    invoice_id = invoices()["accepted"][0]["id"]

    response = client.put(f"/update_invoice/{invoice_id}", json={"total_amount": "2.000,50", "currency": "chf"})
    assert response.status_code == 200
    saved = invoices()["accepted"][0]
    assert saved["total_amount"] == "2000.50"
    assert saved["currency"] == "CHF"

    assert client.put(f"/update_invoice/{invoice_id}", json={"tags": "Food, travel , "}).status_code == 200
    assert invoices()["accepted"][0]["tags"] == "food, travel"

    assert client.put(f"/update_invoice/{invoice_id}", json={"total_amount": ""}).status_code == 400
    assert client.put(f"/update_invoice/{invoice_id}", json={"id = 1, tax_number": "x"}).status_code == 400
    assert client.put("/update_invoice/9999", json={"vat_id": "x"}).status_code == 404


def test_edit_cannot_take_another_invoices_number():
    client.post("/save_invoice", json={"fields": GOOD_FIELDS, "language": "de"})
    client.post("/save_invoice", json={"fields": {**GOOD_FIELDS, "invoice_number": "RE-2"}, "language": "de"})
    second_id = invoices()["accepted"][0]["id"]
    assert client.put(f"/update_invoice/{second_id}", json={"invoice_number": "RE-1"}).status_code == 409
    assert client.put(f"/update_invoice/{second_id}", json={"invoice_number": "RE-2"}).status_code == 200


def test_delete_saved_and_rejected_keeps_history():
    client.post("/save_invoice", json={"fields": GOOD_FIELDS, "language": "de"})
    client.post("/reject_invoice", json={"invoice_number": "RE-9", "reason": "Missing fields: Tax Number"})
    data = invoices()

    assert client.delete(f"/delete_invoice/{data['accepted'][0]['id']}").status_code == 200
    assert client.delete(f"/delete_rejected/{data['rejected'][0]['id']}").status_code == 200
    assert client.delete("/delete_rejected/9999").status_code == 404

    data = invoices()
    assert data["accepted"] == [] and data["rejected"] == []
    assert len(data["history"]) == 2


def exported_rows(**form):
    response = client.post("/export_excel", data={"type": "accepted", **form})
    assert response.status_code == 200
    return list(load_workbook(io.BytesIO(response.content)).active.iter_rows(min_row=2, values_only=True))


def test_export_date_filter():
    client.post("/save_invoice", json={"fields": GOOD_FIELDS, "language": "de"})
    client.post("/save_invoice", json={"fields": {**GOOD_FIELDS, "invoice_number": "RE-2", "date": "2024-01-10"}, "language": "de"})

    assert len(exported_rows()) == 2  # no range: everything
    rows = exported_rows(date_type="issue_date", from_date="2025-01-01")
    assert [r[2] for r in rows] == ["RE-1"]
    assert rows[0][8] == 1785.0 and rows[0][9] == "EUR"
    assert not rows[0][10]  # Tags column is empty: no tags were sent in this test
    assert [r[2] for r in exported_rows(date_type="issue_date", to_date="2024-12-31")] == ["RE-2"]


def test_export_rejects_unknown_date_column():
    response = client.post("/export_excel", data={"type": "accepted", "date_type": "x); DROP TABLE invoices;--"})
    assert response.status_code == 400

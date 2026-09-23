import os
import sqlite3
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

# Always use the DB next to this file, no matter which folder uvicorn is started from
DB_NAME = os.path.join(os.path.dirname(os.path.abspath(__file__)), "invoices.db")

# Column names can't be passed as SQL parameters, so only these are ever put into a query
DATE_COLUMNS = {"processing_date", "issue_date"}
EDITABLE_COLUMNS = {
    "invoice_number", "issue_date", "tax_number", "vat_percent", "vat_amount",
    "vat_id", "total_amount", "exemption_reason",
}

ACCEPTED_COLUMNS = [
    "id", "processing_date", "invoice_number", "issue_date",
    "tax_number", "vat_percent", "vat_amount", "vat_id",
    "total_amount", "exemption_reason", "used_ocr", "language"
]


def init_db():
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()

    cur.execute('''
        CREATE TABLE IF NOT EXISTS invoices (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            processing_date TEXT,
            invoice_number TEXT,
            issue_date TEXT,
            tax_number TEXT,
            vat_percent TEXT,
            vat_amount TEXT,
            vat_id TEXT,
            total_amount TEXT,
            exemption_reason TEXT,
            accepted BOOLEAN,
            reason TEXT,
            used_ocr BOOLEAN,
            language TEXT
        )
    ''')

    # Older databases were created before the language column existed
    columns = [row[1] for row in cur.execute("PRAGMA table_info(invoices)")]
    if "language" not in columns:
        cur.execute("ALTER TABLE invoices ADD COLUMN language TEXT")

    cur.execute('''
        CREATE TABLE IF NOT EXISTS rejected_invoices (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            processing_date TEXT,
            rejection_date TEXT,
            invoice_number TEXT,
            issue_date TEXT,
            reason TEXT
        )
    ''')

    conn.commit()
    conn.close()

def save_invoice(fields, accepted=True, reason="", used_ocr=False, language=""):
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()
    now = datetime.now().isoformat()

    cur.execute('''
        INSERT INTO invoices (
            processing_date, invoice_number, issue_date, tax_number, vat_percent, vat_amount, vat_id,
            total_amount, exemption_reason, accepted, reason, used_ocr, language
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (
        now,
        fields.get("invoice_number", ""),
        fields.get("date", ""),
        fields.get("tax_number", ""),
        fields.get("vat_percent", ""),
        fields.get("vat_amount", ""),
        fields.get("vat_id", ""),
        fields.get("total_amount", ""),
        fields.get("exemption_reason", ""),
        int(accepted),
        reason if not accepted else "",
        int(used_ocr),
        language
    ))
    conn.commit()
    conn.close()

def get_all_invoices():
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()

    accepted = cur.execute(
        f"SELECT {', '.join(ACCEPTED_COLUMNS)} FROM invoices WHERE accepted=1 ORDER BY id DESC"
    ).fetchall()
    rejected = cur.execute(
        "SELECT id, processing_date, rejection_date, invoice_number, issue_date, reason "
        "FROM rejected_invoices ORDER BY id DESC"
    ).fetchall()
    conn.close()

    return {
        "accepted": [dict(zip(ACCEPTED_COLUMNS, row)) for row in accepted],
        "rejected": [
            {
                "id": row[0],
                "processing_date": row[1],
                "rejection_date": row[2],
                "invoice_number": row[3],
                "issue_date": row[4],
                "reason": row[5]
            } for row in rejected
        ]
    }

def save_rejected_invoice(invoice_number: str, issue_date: str, reason: str):
    now = datetime.now().isoformat()
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO rejected_invoices (processing_date, rejection_date, invoice_number, issue_date, reason)
        VALUES (?, ?, ?, ?, ?)
    """, (now, now, invoice_number, issue_date , reason))
    conn.commit()
    conn.close()

def invoice_exists_in_any_table(invoice_number: str) -> bool:
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()

    cur.execute("SELECT COUNT(*) FROM invoices WHERE invoice_number = ?", (invoice_number,))
    count_accepted = cur.fetchone()[0]

    cur.execute("SELECT COUNT(*) FROM rejected_invoices WHERE invoice_number = ?", (invoice_number,))
    count_rejected = cur.fetchone()[0]

    conn.close()
    return count_accepted > 0 or count_rejected > 0

def _to_number(value):
    """Amounts are stored as text like "1234.56"; Excel needs real numbers to sum them."""
    try:
        return float(value)
    except (TypeError, ValueError):
        return value

def export_invoices_to_excel(invoice_type: str, from_date: str, to_date: str, date_type: str = "processing_date"):
    if date_type not in DATE_COLUMNS:
        raise ValueError(f"Invalid date type: {date_type}")

    wb = Workbook()
    sheet = wb.active

    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()

    if invoice_type == "rejected":
        sheet.title = "rejected-invoice"
        headers = ["Index", "Processing Date", "Invoice Number", "Issue Date", "Reason"]
        number_columns = {}
        query = f"""
            SELECT id, processing_date, invoice_number, issue_date, reason
            FROM rejected_invoices
            WHERE DATE({date_type}) BETWEEN ? AND ?
        """
    else:
        sheet.title = "accepted-invoice"
        headers = ["ID", "Processing Date", "Invoice Number", "Issue Date", "Tax Number", "VAT %", "VAT Amount", "VAT ID", "Total", "Exemption Reason"]
        # column index -> Excel number format
        number_columns = {5: "0", 6: "#,##0.00", 8: "#,##0.00"}
        query = f"""
            SELECT id, processing_date, invoice_number, issue_date,  tax_number, vat_percent,
                   vat_amount, vat_id, total_amount, exemption_reason
            FROM invoices
            WHERE accepted = 1 AND DATE({date_type}) BETWEEN ? AND ?
        """
    rows = cur.execute(query, (from_date, to_date)).fetchall()
    conn.close()

    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill(start_color="CDEEEF", end_color="CDEEEF", fill_type="solid")
        sheet.column_dimensions[cell.column_letter].width = 30

    for row in rows:
        row = [_to_number(v) if i in number_columns else v for i, v in enumerate(row)]
        sheet.append(row)
        for i, cell in enumerate(sheet[sheet.max_row]):
            cell.alignment = Alignment(horizontal="center")
            if i in number_columns and isinstance(cell.value, float):
                cell.number_format = number_columns[i]

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()

def update_invoice_by_id(invoice_id: int, updated_fields: dict):
    unknown = set(updated_fields) - EDITABLE_COLUMNS
    if unknown or not updated_fields:
        raise ValueError(f"Invalid fields: {', '.join(sorted(unknown)) or 'none given'}")

    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()
    set_clause = ", ".join(f"{key} = ?" for key in updated_fields.keys())
    values = list(updated_fields.values()) + [invoice_id]

    query = f"UPDATE invoices SET {set_clause} WHERE id = ?"
    cur.execute(query, values)
    conn.commit()
    conn.close()

def delete_invoice_by_id(invoice_id: int):
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()
    cur.execute("DELETE FROM invoices WHERE id = ?", (invoice_id,))
    conn.commit()
    conn.close()
